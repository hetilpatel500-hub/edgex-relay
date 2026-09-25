"""Debt Payoff Planner: snowball vs avalanche vs minimum-only, 10 debts, 25 years."""
import os
import sys
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter as L
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from xlsxkit import Kit, INK, MUTED

ACCENT = "3F6E8C"
K = Kit(ACCENT)
N_DEBTS = 10
MONTHS = 300
FIRST = 8                       # first schedule data row
LAST = FIRST + MONTHS - 1
D0 = 10                         # first debt row on My Debts
D1 = D0 + N_DEBTS - 1
MONEY = '"$"#,##0.00'
MONEY0 = '"$"#,##0'
DATEF = 'mmm yyyy'
SLOT0 = 7                       # first slot column on schedule sheets
SW = 6                          # columns per slot

SAMPLE = [
    ("Visa card", 4200, 24.99, 110),
    ("Car loan", 11500, 6.9, 285),
    ("Store card", 950, 29.99, 35),
    ("Student loan", 18300, 5.5, 200),
    ("Personal loan", 3100, 11.5, 95),
]
SHEETS = {"snow": "Snowball", "ava": "Avalanche", "min": "Minimum Only"}


def debts_sheet(wb, sample):
    ws = wb.active
    ws.title = "My Debts"
    K.sheet_setup(ws, [2, 24, 16, 12, 16, 13, 13, 3, 3, 3, 3], tab=ACCENT)
    K.title(ws, "My Debts", "List every debt once. Everything else in this file updates from this page.", span=6)
    K.section(ws["B5"], "Your plan settings")
    K.label(ws["B6"], "First month of your plan")
    ws["C6"] = dt.date(2026, 10, 1) if sample else dt.date(2026, 10, 1)
    K.body(ws["C6"], "mmm yyyy", inp=True)
    K.label(ws["B7"], "Extra you can pay each month")
    ws["C7"] = 250 if sample else 0
    K.body(ws["C7"], MONEY, inp=True)
    K.label(ws["B8"], "Method you will follow")
    ws["C8"] = "Snowball"
    K.body(ws["C8"], inp=True)
    K.list_validation(ws, "C8", '"Snowball,Avalanche"', "Snowball = smallest balance first. Avalanche = highest interest first.")
    K.note(ws, "D6", "Pick the 1st of the month you start.")
    K.note(ws, "D7", "On top of your minimum payments. Can be $0.")
    K.note(ws, "D8", "Not sure? See the Compare tab.")

    hdr = D0 - 1
    K.header_row(ws, hdr, 2, ["Debt name", "Balance now", "APR (%)", "Minimum payment", "Snowball order", "Avalanche order"])
    for i in range(N_DEBTS):
        r = D0 + i
        if sample and i < len(SAMPLE):
            n, b, a, m = SAMPLE[i]
            ws.cell(row=r, column=2, value=n)
            ws.cell(row=r, column=3, value=b)
            ws.cell(row=r, column=4, value=a)
            ws.cell(row=r, column=5, value=m)
        for c, fmt in ((2, None), (3, MONEY), (4, '0.00'), (5, MONEY)):
            K.body(ws.cell(row=r, column=c), fmt, inp=True)
        # hidden sort keys + full ranks (H..K)
        ws.cell(row=r, column=8, value='=IF(OR(B{r}="",N(C{r})<=0),1E9+ROW(),C{r}+ROW()/1E6)'.format(r=r))
        ws.cell(row=r, column=9, value='=IF(OR(B{r}="",N(C{r})<=0),1E9+ROW(),-N(D{r})+ROW()/1E6)'.format(r=r))
        ws.cell(row=r, column=10, value='=RANK(H{r},$H${a}:$H${b},1)'.format(r=r, a=D0, b=D1))
        ws.cell(row=r, column=11, value='=RANK(I{r},$I${a}:$I${b},1)'.format(r=r, a=D0, b=D1))
        ws.cell(row=r, column=6, value='=IF(OR(B{r}="",N(C{r})<=0),"",J{r})'.format(r=r))
        ws.cell(row=r, column=7, value='=IF(OR(B{r}="",N(C{r})<=0),"",K{r})'.format(r=r))
        for c in (6, 7):
            K.body(ws.cell(row=r, column=c), '0', align="center")
    for col in "HIJK":
        ws.column_dimensions[col].hidden = True
    t = D1 + 1
    K.label(ws.cell(row=t, column=2), "Totals", color=INK)
    ws.cell(row=t, column=3, value="=SUM(C{a}:C{b})".format(a=D0, b=D1))
    ws.cell(row=t, column=5, value="=SUM(E{a}:E{b})".format(a=D0, b=D1))
    for c in (3, 5):
        K.body(ws.cell(row=t, column=c), MONEY, bold=True, band=True)
    K.label(ws.cell(row=t + 1, column=2), "Monthly debt budget", color=INK)
    ws.cell(row=t + 1, column=5, value="=E{t}+C7".format(t=t))
    K.body(ws.cell(row=t + 1, column=5), MONEY, bold=True, band=True)
    K.note(ws, "F%d" % (t + 1), "Minimums + extra. Keep paying this every month.")

    # plan summary
    s = t + 3
    K.section(ws.cell(row=s, column=2), "Your plan at a glance")
    K.kpi(ws, s + 1, 2, "Debt-free by", '=IF($C$8="Avalanche",Compare!C7,Compare!C6)', DATEF, 1)
    K.kpi(ws, s + 1, 3, "Months to go", '=IF($C$8="Avalanche",Compare!D7,Compare!D6)', '0', 2)
    K.kpi(ws, s + 1, 5, "Total interest", '=IF($C$8="Avalanche",Compare!E7,Compare!E6)', MONEY0, 1)
    K.kpi(ws, s + 1, 6, "Interest saved", '=IF($C$8="Avalanche",Compare!G7,Compare!G6)', MONEY0, 2)
    K.note(ws, "B%d" % (s + 4), "Interest saved is compared with paying only the minimums. Estimates only: interest is worked out monthly as balance x APR / 12.", italic=True)
    ws.freeze_panes = "A5"
    return ws


def schedule_sheet(wb, mode):
    ws = wb.create_sheet(SHEETS[mode])
    widths = [7, 11, 14, 13, 12, 3] + [12, 10, 11, 11, 12, 3] * N_DEBTS
    K.sheet_setup(ws, widths, tab=K.light)
    titles = {"snow": ("Snowball schedule", "Smallest balance first. Each paid-off debt's payment rolls into the next one."),
              "ava": ("Avalanche schedule", "Highest APR first. Each paid-off debt's payment rolls into the next one."),
              "min": ("Minimum-only schedule", "What happens if you only ever pay the minimums. Shown for comparison.")}
    K.title(ws, *titles[mode], span=5)
    budget = "'My Debts'!$E$%d" % (D1 + 2)
    K.header_row(ws, 7, 1, ["Month", "Date", "Total owed (end)", "Paid this month", "Interest this month"])
    ws.cell(row=7, column=6).value = None
    for k in range(N_DEBTS):
        b = SLOT0 + k * SW
        S, I, M, X, E, P = (L(b + j) for j in range(6))
        if mode == "snow":
            src = "MATCH({k},'My Debts'!$J${a}:$J${z},0)".format(k=k + 1, a=D0, z=D1)
        elif mode == "ava":
            src = "MATCH({k},'My Debts'!$K${a}:$K${z},0)".format(k=k + 1, a=D0, z=D1)
        else:
            src = str(k + 1)
        # helper row 6 (hidden): src index, APR, min, start balance, name
        ws.cell(row=6, column=b, value="=" + src)
        ws.cell(row=6, column=b + 1, value="=N(INDEX('My Debts'!$D${a}:$D${z},{s}$6))/100".format(a=D0, z=D1, s=S))
        ws.cell(row=6, column=b + 2, value="=IF(INDEX('My Debts'!$B${a}:$B${z},{s}$6)=\"\",0,N(INDEX('My Debts'!$E${a}:$E${z},{s}$6)))".format(a=D0, z=D1, s=S))
        ws.cell(row=6, column=b + 3, value="=IF(INDEX('My Debts'!$B${a}:$B${z},{s}$6)=\"\",0,N(INDEX('My Debts'!$C${a}:$C${z},{s}$6)))".format(a=D0, z=D1, s=S))
        ws.cell(row=6, column=b + 4, value="=IF(INDEX('My Debts'!$B${a}:$B${z},{s}$6)=\"\",\"\",INDEX('My Debts'!$B${a}:$B${z},{s}$6))".format(a=D0, z=D1, s=S))
        # visible slot title row 5
        lab = "Debt {k}".format(k=k + 1) if mode == "min" else "Payoff #{k}".format(k=k + 1)
        ws.cell(row=5, column=b, value='="{lab}: "&IF({e}6="","(empty)",{e}6)'.format(lab=lab, e=E))
        ws.merge_cells(start_row=5, start_column=b, end_row=5, end_column=b + 4)
        c5 = ws.cell(row=5, column=b)
        c5.font = Font(name="Arial", size=10, bold=True, color=ACCENT)
        c5.fill = K.fill(K.lighter)
        c5.alignment = Alignment(horizontal="center")
        K.header_row(ws, 7, b, ["Start balance", "Interest", "Minimum paid", "Extra paid", "End balance"])
        for r in range(FIRST, LAST + 1):
            prev_pool = "$F{r}".format(r=r) if k == 0 else "{p}{r}".format(p=L(b - 1), r=r)
            start = "={x}$6".format(x=X) if r == FIRST else "={e}{p}".format(e=E, p=r - 1)
            ws.cell(row=r, column=b, value=start)
            ws.cell(row=r, column=b + 1, value="=ROUND({s}{r}*{i}$6/12,2)".format(s=S, r=r, i=I))
            ws.cell(row=r, column=b + 2, value="=MIN({m}$6,{s}{r}+{i}{r})".format(m=M, s=S, r=r, i=I))
            ws.cell(row=r, column=b + 3, value="=MAX(0,MIN({s}{r}+{i}{r}-{m}{r},{pp}))".format(s=S, i=I, m=M, r=r, pp=prev_pool))
            ws.cell(row=r, column=b + 4, value="=ROUND({s}{r}+{i}{r}-{m}{r}-{x}{r},2)".format(s=S, i=I, m=M, x=X, r=r))
            ws.cell(row=r, column=b + 5, value="={pp}-{x}{r}".format(pp=prev_pool, x=X, r=r))
            for j in range(5):
                c = ws.cell(row=r, column=b + j)
                c.number_format = MONEY
                c.font = Font(name="Arial", size=9, color=INK)
        ws.column_dimensions[P].hidden = True
    ws.row_dimensions[6].hidden = True
    ends = ["{c}{{r}}".format(c=L(SLOT0 + k * SW + 4)) for k in range(N_DEBTS)]
    mins = ["{c}{{r}}".format(c=L(SLOT0 + k * SW + 2)) for k in range(N_DEBTS)]
    xtra = ["{c}{{r}}".format(c=L(SLOT0 + k * SW + 3)) for k in range(N_DEBTS)]
    ints = ["{c}{{r}}".format(c=L(SLOT0 + k * SW + 1)) for k in range(N_DEBTS)]
    for r in range(FIRST, LAST + 1):
        ws.cell(row=r, column=1, value=r - FIRST + 1)
        ws.cell(row=r, column=2, value="=EDATE('My Debts'!$C$6,A{r}-1)".format(r=r))
        ws.cell(row=r, column=3, value="=" + "+".join(e.format(r=r) for e in ends))
        ws.cell(row=r, column=4, value="=" + "+".join(e.format(r=r) for e in mins + xtra))
        ws.cell(row=r, column=5, value="=" + "+".join(e.format(r=r) for e in ints))
        if mode == "min":
            ws.cell(row=r, column=6, value=0)
        else:
            ws.cell(row=r, column=6, value="=MAX(0,{b}-({m}))".format(b=budget, m="+".join(e.format(r=r) for e in mins)))
        for c, fmt in ((1, "0"), (2, DATEF), (3, MONEY), (4, MONEY), (5, MONEY)):
            cell = ws.cell(row=r, column=c)
            cell.number_format = fmt
            cell.font = Font(name="Arial", size=9, bold=(c == 3), color=INK)
            if r % 2:
                cell.fill = K.fill(K.lighter)
    ws.column_dimensions["F"].hidden = True
    ws.freeze_panes = ws.cell(row=FIRST, column=3)
    # grey out months after debt-free
    ws.conditional_formatting.add("A{a}:E{z}".format(a=FIRST, z=LAST),
                                  FormulaRule(formula=["$C{a}<=0.005".format(a=FIRST)], font=Font(color="B8BEC4")))
    return ws


def compare_sheet(wb):
    ws = wb.create_sheet("Compare", 1)
    K.sheet_setup(ws, [2, 18, 14, 11, 15, 15, 16, 3, 10, 22, 14, 14], tab=ACCENT)
    K.title(ws, "Compare methods", "Same monthly budget, three ways to use it. Pick yours on the My Debts tab.", span=10)
    K.header_row(ws, 5, 2, ["Method", "Debt-free by", "Months", "Total interest", "Total paid", "Interest saved vs minimums"])
    rows = [("Snowball", "Snowball"), ("Avalanche", "Avalanche"), ("Minimums only", "Minimum Only")]
    for i, (label, sh) in enumerate(rows):
        r = 6 + i
        q = "'%s'" % sh
        ws.cell(row=r, column=2, value=label)
        ws.cell(row=r, column=4, value='=IF(\'My Debts\'!$C${t}<=0,0,IF(COUNTIF({q}!$C${a}:$C${z},">0.005")>={n},"25+ yrs",COUNTIF({q}!$C${a}:$C${z},">0.005")+1))'.format(q=q, a=FIRST, z=LAST, n=MONTHS, t=D1 + 1))
        ws.cell(row=r, column=3, value='=IF(ISNUMBER(D{r}),EDATE(\'My Debts\'!$C$6,MAX(D{r},1)-1),"Not in 25 yrs")'.format(r=r))
        ws.cell(row=r, column=5, value="=SUM({q}!$E${a}:$E${z})".format(q=q, a=FIRST, z=LAST))
        ws.cell(row=r, column=6, value="=SUM({q}!$D${a}:$D${z})".format(q=q, a=FIRST, z=LAST))
        ws.cell(row=r, column=7, value="=E8-E{r}".format(r=r) if i < 2 else "—")
        for c, fmt in ((2, None), (3, DATEF), (4, "0"), (5, MONEY), (6, MONEY), (7, MONEY)):
            K.body(ws.cell(row=r, column=c), fmt, bold=(c == 2), band=(i == 2))
    K.note(ws, "B10", "If minimums never clear a debt within 25 years, the minimum-only interest shown is only the first 25 years, so real savings are larger.", italic=True)
    K.note(ws, "B11", "Snowball often feels better (quick wins). Avalanche usually costs less. The best plan is the one you stick with.", italic=True)

    for j, (label, sh, col) in enumerate((("Snowball order", "Snowball", 2), ("Avalanche order", "Avalanche", 9))):
        top = 13
        K.section(ws.cell(row=top, column=col), label)
        K.header_row(ws, top + 1, col, ["#", "Debt", "Paid off", "Interest paid"] if col == 9 else ["Debt", "Paid off", "Interest paid"])
    # snowball table B..D, avalanche table I..L
    for k in range(N_DEBTS):
        r = 15 + k
        base = SLOT0 + k * SW
        for sh, c0, with_num in (("Snowball", 2, False), ("Avalanche", 10, True)):
            q = "'%s'" % sh
            endc, intc, namec = L(base + 4), L(base + 1), L(base + 4)
            name = "={q}!{n}6".format(q=q, n=namec)
            months = 'COUNTIF({q}!${e}${a}:${e}${z},">0.005")'.format(q=q, e=endc, a=FIRST, z=LAST)
            paid = '=IF({q}!{n}6="","",IF({m}>={N},"Not in 25 yrs",EDATE(\'My Debts\'!$C$6,{m})))'.format(q=q, n=namec, m=months, N=MONTHS)
            intr = '=IF({q}!{n}6="","",SUM({q}!${i}${a}:${i}${z}))'.format(q=q, n=namec, i=intc, a=FIRST, z=LAST)
            if with_num:
                ws.cell(row=r, column=c0 - 1, value=k + 1)
                K.body(ws.cell(row=r, column=c0 - 1), "0", align="center", band=bool(k % 2))
            ws.cell(row=r, column=c0, value=name)
            ws.cell(row=r, column=c0 + 1, value=paid)
            ws.cell(row=r, column=c0 + 2, value=intr)
            for c, fmt in ((c0, None), (c0 + 1, DATEF), (c0 + 2, MONEY)):
                K.body(ws.cell(row=r, column=c), fmt, band=bool(k % 2))
    # fix: paid-off date is the month the balance hits zero => EDATE(start, months) where months = count of positive months
    # line chart of total owed
    ch = LineChart()
    ch.title = "Total owed over time"
    ch.height, ch.width = 8.5, 22
    ch.y_axis.title = "Total owed"
    ch.x_axis.title = "Month"
    ch.y_axis.numFmt = '"$"#,##0'
    for sh in ("Snowball", "Avalanche", "Minimum Only"):
        data = Reference(wb[sh], min_col=3, min_row=FIRST, max_row=FIRST + 179)
        ch.add_data(data, titles_from_data=False)
    for s, name in zip(ch.series, ("Snowball", "Avalanche", "Minimums only")):
        from openpyxl.chart.series import SeriesLabel
        s.tx = SeriesLabel(v=name)
        s.smooth = False
    ch.series[0].graphicalProperties.line.solidFill = ACCENT
    ch.series[1].graphicalProperties.line.solidFill = "C7823F"
    ch.series[2].graphicalProperties.line.solidFill = "A9B1B8"
    ch.x_axis.tickLblSkip = 12
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ws.add_chart(ch, "B27")
    K.note(ws, "B26", "Chart shows the first 15 years.", italic=True)
    return ws


def tracker_sheet(wb):
    ws = wb.create_sheet("Progress Log")
    K.sheet_setup(ws, [2, 14, 16, 16, 14, 36], tab=K.light, landscape=False)
    K.title(ws, "Progress log", "Once a month, write down what you really owe. Watch it shrink.", span=5)
    K.label(ws["B5"], "Starting total")
    ws["C5"] = "='My Debts'!C%d" % (D1 + 1)
    K.body(ws["C5"], MONEY, bold=True, band=True)
    K.header_row(ws, 7, 2, ["Month", "Total you owe now", "Paid this month", "Progress", "Notes / wins"])
    for i in range(60):
        r = 8 + i
        ws.cell(row=r, column=2, value="=EDATE('My Debts'!$C$6,%d)" % i)
        K.body(ws.cell(row=r, column=2), DATEF)
        for c, fmt in ((3, MONEY), (4, MONEY), (6, None)):
            K.body(ws.cell(row=r, column=c), fmt, inp=True)
        ws.cell(row=r, column=5, value='=IF(OR(C{r}="",$C$5<=0),"",1-C{r}/$C$5)'.format(r=r))
        K.body(ws.cell(row=r, column=5), "0%", align="center")
    from openpyxl.formatting.rule import DataBarRule
    ws.conditional_formatting.add("E8:E67", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=ACCENT))
    ws.freeze_panes = "A8"


def start_sheet(wb):
    ws = wb.create_sheet("Start Here", 0)
    K.start_here(ws, "Debt Payoff Planner",
                 "See exactly when you will be debt-free, and compare the snowball and avalanche methods side by side. "
                 "Works in Microsoft Excel and Google Sheets.",
                 ["Go to My Debts. Type each debt's name, balance, APR and minimum payment (up to 10 debts).",
                  "Enter the first month of your plan and any extra amount you can pay each month on top of the minimums.",
                  "Open Compare to see your debt-free date, total interest and payoff order for each method.",
                  "Pick Snowball or Avalanche on My Debts. Pay your monthly debt budget every month, following the payoff order.",
                  "Each month, log what you really owe on Progress Log.",
                  "Want the detail? The Snowball, Avalanche and Minimum Only tabs show every month for up to 25 years."],
                 ["Google Sheets: File > Import > Upload, or upload to Google Drive and open with Google Sheets.",
                  "Type APR as a plain number, e.g. 22.99 for 22.99%.",
                  "Interest is estimated monthly (balance x APR / 12). Your lender may calculate it daily, so their numbers will be a little different.",
                  "This is a planning tool, not financial advice. Check payoff amounts with your lender.",
                  "Please do not change the white formula cells. If something breaks, re-download your original file from Etsy (Purchases > Download files)."])


def build(path, sample=False):
    wb = Workbook()
    debts_sheet(wb, sample)
    for m in ("snow", "ava", "min"):
        schedule_sheet(wb, m)
    compare_sheet(wb)
    tracker_sheet(wb)
    start_sheet(wb)
    if sample:
        for _n, _a in {'Start Here': 'A1:D22', 'My Debts': 'A1:G30', 'Compare': 'A1:L45', 'Snowball': 'A1:Q30', 'Avalanche': 'A1:Q30', 'Minimum Only': 'A1:Q30', 'Progress Log': 'A1:F24'}.items():
            wb[_n].print_area = _a
    wb.active = 0
    wb.save(path)
    print(path)


if __name__ == "__main__":
    here = os.path.dirname(os.path.abspath(__file__))
    out = os.path.join(here, "..", "03-debt-payoff-planner")
    os.makedirs(out, exist_ok=True)
    build(os.path.join(out, "Debt-Payoff-Planner.xlsx"))
    samp = sys.argv[1] if len(sys.argv) > 1 else os.path.join(here, "samples")
    os.makedirs(samp, exist_ok=True)
    build(os.path.join(samp, "Debt-Payoff-Planner_SAMPLE.xlsx"), sample=True)
