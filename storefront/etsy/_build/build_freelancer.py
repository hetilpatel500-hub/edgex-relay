"""Freelancer Income, Expense & Tax Set-Aside Tracker."""
import os
import sys
import random
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule, DataBarRule
from xlsxkit import Kit, INK, MUTED

ACCENT = "2F7D6D"
K = Kit(ACCENT)
MONEY = '"$"#,##0.00'
MONEY0 = '"$"#,##0'
DATEF = 'mmm d, yyyy'
ROWS = 500
I0, I1 = 7, 7 + ROWS - 1          # income / expense data rows
V0, V1 = 7, 7 + 200 - 1           # invoice rows
C0, C1 = 14, 43                   # clients on Setup
E0, E1 = 14, 33                   # categories on Setup
YEAR = "Setup!$C$6"
RATE = "Setup!$C$8"

CATS = [("Advertising & marketing", 1), ("Software & subscriptions", 1), ("Office supplies", 1),
        ("Equipment & computers", 1), ("Phone & internet", 1), ("Professional services", 1),
        ("Contractors", 1), ("Travel", 1), ("Business meals", 0.5), ("Education & training", 1),
        ("Bank & payment fees", 1), ("Business insurance", 1), ("Rent / coworking", 1), ("Home office", 1),
        ("Vehicle / mileage", 1), ("Shipping & postage", 1), ("Licenses & dues", 1), ("Other", 1)]
CLIENTS = ["Northwind Studio", "Bright Path Co.", "Maple & Oak", "Harbor Health", "Cedar Labs", "Blue Fern Media"]


def setup_sheet(wb, sample):
    ws = wb.create_sheet("Setup")
    K.sheet_setup(ws, [2, 30, 16, 4, 28, 14, 40], tab=ACCENT, landscape=False)
    K.title(ws, "Setup", "Fill this in once. Every other tab reads from here.", span=6)
    K.label(ws["B6"], "Tax year")
    ws["C6"] = 2026
    K.body(ws["C6"], "0", inp=True, align="center")
    K.label(ws["B7"], "Business name")
    ws["C7"] = "My Freelance Business"
    K.body(ws["C7"], inp=True)
    ws.merge_cells("C7:E7")
    K.label(ws["B8"], "Tax set-aside %")
    ws["C8"] = 0.25
    K.body(ws["C8"], "0%", inp=True, align="center")
    K.note(ws, "E8", "The share of income minus deductible expenses you move to savings for tax. Ask a tax professional what fits you.")
    K.label(ws["B9"], "Monthly income goal")
    ws["C9"] = 6000 if sample else 0
    K.body(ws["C9"], MONEY0, inp=True)
    K.header_row(ws, 13, 2, ["Your clients"])
    K.header_row(ws, 13, 5, ["Expense categories", "Deductible %"])
    for i in range(C1 - C0 + 1):
        c = ws.cell(row=C0 + i, column=2, value=(CLIENTS[i] if (sample and i < len(CLIENTS)) else None))
        K.body(c, inp=True)
    for i in range(E1 - E0 + 1):
        name, pct = CATS[i] if i < len(CATS) else (None, None)
        K.body(ws.cell(row=E0 + i, column=5, value=name), inp=True)
        K.body(ws.cell(row=E0 + i, column=6, value=pct), "0%", inp=True, align="center")
    K.note(ws, "E35", "Deductible % is the share of that cost you can claim. Rules differ by country;", wrap=False)
    K.note(ws, "E36", "for example, US business meals are usually 50% deductible. Check with a tax pro.")
    return ws


def income_sheet(wb, sample):
    ws = wb.create_sheet("Income")
    K.sheet_setup(ws, [2, 14, 24, 34, 12, 14, 18, 30], tab=K.light)
    K.title(ws, "Income", "Log money when it lands in your account. One row per payment.", span=7)
    ws["B5"] = '="Total received in "&Setup!$C$6&":"'
    ws["B5"].font = Font(name="Arial", size=10, bold=True, color=MUTED)
    ws["E5"] = '=SUMIFS(F{a}:F{b},B{a}:B{b},">="&DATE({y},1,1),B{a}:B{b},"<"&DATE({y}+1,1,1))'.format(a=I0, b=I1, y=YEAR)
    K.body(ws["E5"], MONEY, bold=True, band=True)
    ws.merge_cells("E5:F5")
    K.header_row(ws, 6, 2, ["Date received", "Client", "What it was for", "Invoice #", "Amount", "Paid by", "Notes"])
    K.list_validation(ws, "C{a}:C{b}".format(a=I0, b=I1), "=Setup!$B${a}:$B${b}".format(a=C0, b=C1), "Add clients on the Setup tab.")
    K.list_validation(ws, "G{a}:G{b}".format(a=I0, b=I1), '"Bank transfer,Card,Payment app,Check,Cash,Other"')
    rnd = random.Random(7)
    for i in range(ROWS):
        r = I0 + i
        for c, fmt in ((2, DATEF), (3, None), (4, None), (5, None), (6, MONEY), (7, None), (8, None)):
            K.body(ws.cell(row=r, column=c), fmt, inp=True)
    if sample:
        r = I0
        for m in range(1, 10):
            for _ in range(rnd.randint(2, 4)):
                d = dt.date(2026, m, rnd.randint(1, 28))
                cl = rnd.choice(CLIENTS)
                ws.cell(row=r, column=2, value=d)
                ws.cell(row=r, column=3, value=cl)
                ws.cell(row=r, column=4, value=rnd.choice(["Website redesign", "Monthly retainer", "Brand guide", "Blog posts (4)", "Consulting call", "Landing page"]))
                ws.cell(row=r, column=5, value="INV-%03d" % (r - I0 + 101))
                ws.cell(row=r, column=6, value=rnd.choice([450, 800, 1200, 1500, 2200, 2750, 900, 1850]))
                ws.cell(row=r, column=7, value=rnd.choice(["Bank transfer", "Card", "Payment app"]))
                r += 1
    ws.freeze_panes = "A7"
    return ws


def expense_sheet(wb, sample):
    ws = wb.create_sheet("Expenses")
    K.sheet_setup(ws, [2, 14, 22, 24, 30, 12, 12, 14, 12], tab=K.light)
    K.title(ws, "Expenses", "Every business cost. Pick a category and the deductible part fills in.", span=8)
    ws["B5"] = '="Deductible total in "&Setup!$C$6&":"'
    ws["B5"].font = Font(name="Arial", size=10, bold=True, color=MUTED)
    ws["E5"] = '=SUMIFS(H{a}:H{b},B{a}:B{b},">="&DATE({y},1,1),B{a}:B{b},"<"&DATE({y}+1,1,1))'.format(a=I0, b=I1, y=YEAR)
    K.body(ws["E5"], MONEY, bold=True, band=True)
    K.header_row(ws, 6, 2, ["Date", "Paid to", "Category", "What it was", "Amount", "Deductible %", "Deductible amount", "Receipt saved?"])
    K.list_validation(ws, "D{a}:D{b}".format(a=I0, b=I1), "=Setup!$E${a}:$E${b}".format(a=E0, b=E1), "Edit categories on the Setup tab.")
    K.list_validation(ws, "I{a}:I{b}".format(a=I0, b=I1), '"Yes,No"')
    for i in range(ROWS):
        r = I0 + i
        for c, fmt in ((2, DATEF), (3, None), (4, None), (5, None), (6, MONEY), (9, None)):
            K.body(ws.cell(row=r, column=c), fmt, inp=True, align="center" if c == 9 else None)
        ws.cell(row=r, column=7, value='=IF(D{r}="","",IFERROR(N(INDEX(Setup!$F${a}:$F${b},MATCH(D{r},Setup!$E${a}:$E${b},0))),0))'.format(r=r, a=E0, b=E1))
        ws.cell(row=r, column=8, value='=IF(OR(F{r}="",G{r}=""),"",F{r}*G{r})'.format(r=r))
        K.body(ws.cell(row=r, column=7), "0%", align="center")
        K.body(ws.cell(row=r, column=8), MONEY)
    ws.conditional_formatting.add("I{a}:I{b}".format(a=I0, b=I1),
                                  FormulaRule(formula=['AND($F{a}<>"",$I{a}<>"Yes")'.format(a=I0)], fill=K.fill("FBE3DC")))
    if sample:
        rnd = random.Random(11)
        r = I0
        items = [("Design app", "Software & subscriptions", 54.99), ("Phone carrier", "Phone & internet", 85),
                 ("Coworking space", "Rent / coworking", 220), ("Online course", "Education & training", 149),
                 ("Client lunch", "Business meals", 62.4), ("Payment processor", "Bank & payment fees", 38.2),
                 ("Office store", "Office supplies", 27.9), ("Laptop stand", "Equipment & computers", 79),
                 ("Ads", "Advertising & marketing", 120), ("Accountant", "Professional services", 300)]
        for m in range(1, 10):
            for _ in range(rnd.randint(3, 5)):
                v, cat, amt = rnd.choice(items)
                ws.cell(row=r, column=2, value=dt.date(2026, m, rnd.randint(1, 28)))
                ws.cell(row=r, column=3, value=v)
                ws.cell(row=r, column=4, value=cat)
                ws.cell(row=r, column=5, value=cat.split(" ")[0] + " expense")
                ws.cell(row=r, column=6, value=amt)
                ws.cell(row=r, column=9, value=rnd.choice(["Yes", "Yes", "Yes", "No"]))
                r += 1
        # one big purchase so the sample shows how a loss month is handled
        for c, v in ((2, dt.date(2026, 2, 12)), (3, "Computer store"), (4, "Equipment & computers"),
                     (5, "New laptop and monitor"), (6, 3400), (9, "Yes")):
            ws.cell(row=r, column=c, value=v)
    ws.freeze_panes = "A7"
    return ws


def invoice_sheet(wb, sample):
    ws = wb.create_sheet("Invoices")
    K.sheet_setup(ws, [2, 12, 24, 14, 14, 13, 10, 14, 16, 30], tab=K.light)
    K.title(ws, "Invoices", "Track what you have billed and who still owes you.", span=9)
    K.label(ws["B5"], "Still unpaid:")
    ws["D5"] = '=SUMIFS(F{a}:F{b},G{a}:G{b},"<>Yes")'.format(a=V0, b=V1)
    K.body(ws["D5"], MONEY, bold=True, band=True)
    K.label(ws["F5"], "Overdue:")
    ws["G5"] = '=COUNTIF(I{a}:I{b},"Overdue*")'.format(a=V0, b=V1)
    K.body(ws["G5"], "0", bold=True, band=True, align="center")
    K.header_row(ws, 6, 2, ["Invoice #", "Client", "Date sent", "Due date", "Amount", "Paid?", "Date paid", "Status", "Notes"])
    K.list_validation(ws, "C{a}:C{b}".format(a=V0, b=V1), "=Setup!$B${a}:$B${b}".format(a=C0, b=C1))
    K.list_validation(ws, "G{a}:G{b}".format(a=V0, b=V1), '"Yes,No"')
    for i in range(V1 - V0 + 1):
        r = V0 + i
        for c, fmt in ((2, None), (3, None), (4, DATEF), (5, DATEF), (6, MONEY), (7, None), (8, DATEF), (10, None)):
            K.body(ws.cell(row=r, column=c), fmt, inp=True, align="center" if c == 7 else None)
        ws.cell(row=r, column=9, value=('=IF(F{r}="","",IF(G{r}="Yes","Paid",IF(E{r}="","Unpaid",'
                                        'IF(TODAY()>E{r},"Overdue "&(TODAY()-E{r})&" days","Due in "&(E{r}-TODAY())&" days"))))').format(r=r))
        K.body(ws.cell(row=r, column=9), None, align="center")
    ws.conditional_formatting.add("I{a}:I{b}".format(a=V0, b=V1),
                                  FormulaRule(formula=['LEFT($I{a},7)="Overdue"'.format(a=V0)], fill=K.fill("FBE3DC"), font=Font(color="A33A2B", bold=True)))
    ws.conditional_formatting.add("I{a}:I{b}".format(a=V0, b=V1),
                                  FormulaRule(formula=['$I{a}="Paid"'.format(a=V0)], font=Font(color=ACCENT, bold=True)))
    if sample:
        today = dt.date(2026, 9, 25)
        rows = [("INV-140", CLIENTS[0], -40, 1800, "Yes"), ("INV-141", CLIENTS[2], -30, 950, "Yes"),
                ("INV-142", CLIENTS[1], -45, 2200, "No"), ("INV-143", CLIENTS[3], -12, 1200, "No"),
                ("INV-144", CLIENTS[4], -3, 2750, "No")]
        for i, (n, cl, off, amt, paid) in enumerate(rows):
            r = V0 + i
            sent = today + dt.timedelta(days=off)
            ws.cell(row=r, column=2, value=n)
            ws.cell(row=r, column=3, value=cl)
            ws.cell(row=r, column=4, value=sent)
            ws.cell(row=r, column=5, value=sent + dt.timedelta(days=30))
            ws.cell(row=r, column=6, value=amt)
            ws.cell(row=r, column=7, value=paid)
            if paid == "Yes":
                ws.cell(row=r, column=8, value=sent + dt.timedelta(days=21))
    ws.freeze_panes = "A7"
    return ws


def monthly_sheet(wb):
    """Set-aside = rate x year-to-date (income - deductible expenses), never below 0.
    Monthly and quarterly amounts are the change in that running figure, so a loss
    month shows a negative amount and every total reconciles to rate x the year's figure."""
    ws = wb.create_sheet("Monthly & Tax")
    K.sheet_setup(ws, [2, 12, 13, 13, 14, 15, 15, 15, 15, 13, 11], tab=ACCENT)
    K.title(ws, "Monthly summary & tax set-aside", None, span=10)
    ws["B3"] = '=Setup!$C$7&"  ·  "&Setup!$C$6&"  ·  setting aside "&TEXT(Setup!$C$8,"0%")&" of income minus deductible expenses"'
    ws["B3"].font = Font(name="Arial", size=10, color=MUTED)
    K.header_row(ws, 6, 2, ["Month", "Income", "All expenses", "Deductible expenses", "Income minus deductible",
                             "Year so far (income minus deductible)", "Set aside this month", "Set aside so far", "Income goal", "vs goal"], height=42)
    inc = "Income!$F${a}:$F${b}".format(a=I0, b=I1)
    ind = "Income!$B${a}:$B${b}".format(a=I0, b=I1)
    exa = "Expenses!$F${a}:$F${b}".format(a=I0, b=I1)
    exd = "Expenses!$H${a}:$H${b}".format(a=I0, b=I1)
    exdt = "Expenses!$B${a}:$B${b}".format(a=I0, b=I1)
    for m in range(1, 13):
        r = 6 + m
        lo = 'DATE({y},{m},1)'.format(y=YEAR, m=m)
        hi = 'DATE({y},{m},1)'.format(y=YEAR, m=m + 1)
        ws.cell(row=r, column=2, value="=" + lo)
        ws.cell(row=r, column=3, value='=SUMIFS({v},{d},">="&{lo},{d},"<"&{hi})'.format(v=inc, d=ind, lo=lo, hi=hi))
        ws.cell(row=r, column=4, value='=SUMIFS({v},{d},">="&{lo},{d},"<"&{hi})'.format(v=exa, d=exdt, lo=lo, hi=hi))
        ws.cell(row=r, column=5, value='=SUMIFS({v},{d},">="&{lo},{d},"<"&{hi})'.format(v=exd, d=exdt, lo=lo, hi=hi))
        ws.cell(row=r, column=6, value="=C{r}-E{r}".format(r=r))
        ws.cell(row=r, column=7, value="=SUM($F$7:F{r})".format(r=r))
        ws.cell(row=r, column=9, value="=MAX(0,G{r})*{rate}".format(r=r, rate=RATE))
        ws.cell(row=r, column=8, value=("=I7" if m == 1 else "=I{r}-I{p}".format(r=r, p=r - 1)))
        ws.cell(row=r, column=10, value="=Setup!$C$9")
        ws.cell(row=r, column=11, value='=IF(J{r}>0,C{r}/J{r},"")'.format(r=r))
        band = bool(m % 2 == 0)
        for c, fmt in ((2, "mmmm"), (3, MONEY), (4, MONEY), (5, MONEY), (6, MONEY), (7, MONEY), (8, MONEY), (9, MONEY), (10, MONEY0), (11, "0%")):
            K.body(ws.cell(row=r, column=c), fmt, band=band, bold=(c in (2, 8)))
    K.label(ws["B19"], "Year total", color=INK)
    totals = {3: "=SUM(C7:C18)", 4: "=SUM(D7:D18)", 5: "=SUM(E7:E18)", 6: "=SUM(F7:F18)", 7: "=G18", 8: "=SUM(H7:H18)", 9: "=I18"}
    for c, f in totals.items():
        ws.cell(row=19, column=c, value=f)
        K.body(ws.cell(row=19, column=c), MONEY, bold=True, band=True)
    ws.conditional_formatting.add("F7:H19", FormulaRule(formula=["F7<0"], font=Font(color="A33A2B")))
    ws.conditional_formatting.add("K7:K18", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1.5, color=ACCENT))

    K.section(ws["B22"], "Quarterly tax set-aside")
    K.header_row(ws, 23, 2, ["Quarter", "Income", "Deductible expenses", "Income minus deductible", "Set aside for this quarter",
                              "Due date", "Amount paid", "Date paid", "Set aside so far minus paid so far"], height=42)
    dues = ["=DATE({y},4,15)", "=DATE({y},6,15)", "=DATE({y},9,15)", "=DATE({y}+1,1,15)"]
    # US federal estimated tax periods are uneven: Jan-Mar, Apr-May, Jun-Aug, Sep-Dec
    rngs = [(7, 9), (10, 11), (12, 14), (15, 18)]
    for q in range(4):
        r = 24 + q
        a, b = rngs[q]
        ws.cell(row=r, column=2, value="Q%d (%s)" % (q + 1, ["Jan-Mar", "Apr-May", "Jun-Aug", "Sep-Dec"][q]))
        ws.cell(row=r, column=3, value="=SUM(C{a}:C{b})".format(a=a, b=b))
        ws.cell(row=r, column=4, value="=SUM(E{a}:E{b})".format(a=a, b=b))
        ws.cell(row=r, column=5, value="=C{r}-D{r}".format(r=r))
        ws.cell(row=r, column=6, value=("=I{b}".format(b=b) if q == 0 else "=I{b}-I{p}".format(b=b, p=rngs[q - 1][1])))
        ws.cell(row=r, column=7, value=dues[q].format(y=YEAR))
        ws.cell(row=r, column=10, value="=I{b}-SUM($H$24:H{r})".format(b=b, r=r))
        for c, fmt, inp in ((2, None, False), (3, MONEY, False), (4, MONEY, False), (5, MONEY, False), (6, MONEY, False),
                            (7, DATEF, True), (8, MONEY, True), (9, DATEF, True), (10, MONEY, False)):
            K.body(ws.cell(row=r, column=c), fmt, inp=inp, bold=(c == 6))
    K.label(ws["B28"], "Year total", color=INK)
    for c, f in {3: "=SUM(C24:C27)", 4: "=SUM(D24:D27)", 5: "=SUM(E24:E27)", 6: "=SUM(F24:F27)", 8: "=SUM(H24:H27)", 10: "=J27"}.items():
        ws.cell(row=28, column=c, value=f)
        K.body(ws.cell(row=28, column=c), MONEY, bold=True, band=True)
    ws.conditional_formatting.add("E24:F28", FormulaRule(formula=["E24<0"], font=Font(color="A33A2B")))
    notes = [
        "How set-aside works: it is always your set-aside % of the year so far (income minus deductible expenses), never below zero. "
        "A loss month or quarter lowers that figure, so its set-aside shows as a negative amount: money you already put aside that you no longer need to add. "
        "The year total always equals your % of the full year's income minus deductible expenses.",
        "Income minus deductible expenses is not your final taxable profit: things like self-employment tax, allowances and other income are not included. "
        "Quarters and due dates follow the US federal estimated-tax schedule (weekends and holidays can move them); outside the US, type your own. "
        "This is a savings guide, not a tax calculation.",
    ]
    for i, n in enumerate(notes):
        r = 30 + i * 3
        ws.merge_cells(start_row=r, start_column=2, end_row=r + 2, end_column=11)
        K.note(ws, "B%d" % r, n, italic=True, wrap=True)
        for rr in range(r, r + 3):
            ws.row_dimensions[rr].height = 22
    ws.freeze_panes = "A7"
    return ws


def breakdown_sheet(wb):
    ws = wb.create_sheet("Breakdown")
    K.sheet_setup(ws, [2, 28, 15, 15, 12, 3, 26, 15, 12, 12], tab=K.light)
    K.title(ws, "Where it comes from, where it goes", None, span=9)
    ws["B3"] = '="Totals for "&Setup!$C$6'
    ws["B3"].font = Font(name="Arial", size=10, color=MUTED)
    K.header_row(ws, 6, 2, ["Expense category", "Spent", "Deductible", "% of spend"])
    K.header_row(ws, 6, 7, ["Client", "Income", "Payments", "% of income"])
    y = YEAR
    for i in range(E1 - E0 + 1):
        r = 7 + i
        ws.cell(row=r, column=2, value='=IF(Setup!E{s}="","",Setup!E{s})'.format(s=E0 + i))
        crit = 'Expenses!$D${a}:$D${b},$B{r},Expenses!$B${a}:$B${b},">="&DATE({y},1,1),Expenses!$B${a}:$B${b},"<"&DATE({y}+1,1,1)'.format(a=I0, b=I1, r=r, y=y)
        ws.cell(row=r, column=3, value='=IF($B{r}="","",SUMIFS(Expenses!$F${a}:$F${b},{c}))'.format(r=r, a=I0, b=I1, c=crit))
        ws.cell(row=r, column=4, value='=IF($B{r}="","",SUMIFS(Expenses!$H${a}:$H${b},{c}))'.format(r=r, a=I0, b=I1, c=crit))
        ws.cell(row=r, column=5, value='=IF(OR($B{r}="",SUM($C$7:$C$26)=0),"",C{r}/SUM($C$7:$C$26))'.format(r=r))
        for c, fmt in ((2, None), (3, MONEY), (4, MONEY), (5, "0%")):
            K.body(ws.cell(row=r, column=c), fmt, band=bool(i % 2))
    for i in range(C1 - C0 + 1):
        r = 7 + i
        ws.cell(row=r, column=7, value='=IF(Setup!B{s}="","",Setup!B{s})'.format(s=C0 + i))
        crit = 'Income!$C${a}:$C${b},$G{r},Income!$B${a}:$B${b},">="&DATE({y},1,1),Income!$B${a}:$B${b},"<"&DATE({y}+1,1,1)'.format(a=I0, b=I1, r=r, y=y)
        ws.cell(row=r, column=8, value='=IF($G{r}="","",SUMIFS(Income!$F${a}:$F${b},{c}))'.format(r=r, a=I0, b=I1, c=crit))
        ws.cell(row=r, column=9, value='=IF($G{r}="","",COUNTIFS({c}))'.format(r=r, c=crit))
        ws.cell(row=r, column=10, value='=IF(OR($G{r}="",SUM($H$7:$H$36)=0),"",H{r}/SUM($H$7:$H$36))'.format(r=r))
        for c, fmt in ((7, None), (8, MONEY), (9, "0"), (10, "0%")):
            K.body(ws.cell(row=r, column=c), fmt, band=bool(i % 2))
    ws.conditional_formatting.add("E7:E26", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=ACCENT))
    ws.conditional_formatting.add("J7:J36", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=ACCENT))
    return ws


def dashboard_sheet(wb):
    ws = wb.create_sheet("Dashboard", 1)
    K.sheet_setup(ws, [2] + [13] * 12, tab=ACCENT)
    K.title(ws, "Dashboard", None, span=12)
    ws["B3"] = '=Setup!$C$7&"  ·  "&Setup!$C$6'
    ws["B3"].font = Font(name="Arial", size=10, color=MUTED)
    m = "'Monthly & Tax'"
    K.kpi(ws, 6, 2, "Income", "={m}!C19".format(m=m), MONEY0, 3)
    K.kpi(ws, 6, 5, "Deductible expenses", "={m}!E19".format(m=m), MONEY0, 3)
    K.kpi(ws, 6, 8, "Income minus deductible", "={m}!F19".format(m=m), MONEY0, 3)
    K.kpi(ws, 6, 11, "All expenses", "={m}!D19".format(m=m), MONEY0, 2)
    K.kpi(ws, 9, 2, "Set aside for tax", "={m}!I19".format(m=m), MONEY0, 3)
    K.kpi(ws, 9, 5, "Tax paid so far", "=SUM({m}!H24:H27)".format(m=m), MONEY0, 3)
    K.kpi(ws, 9, 8, "Unpaid invoices", "=Invoices!D5", MONEY0, 3)
    K.kpi(ws, 9, 11, "Overdue invoices", "=Invoices!G5", "0", 2)
    ch = BarChart()
    ch.type = "col"
    ch.title = "Income vs deductible expenses by month"
    ch.height, ch.width = 8, 26
    mws = wb["Monthly & Tax"]
    ch.add_data(Reference(mws, min_col=3, min_row=6, max_row=18), titles_from_data=True)
    ch.add_data(Reference(mws, min_col=5, min_row=6, max_row=18), titles_from_data=True)
    ch.set_categories(Reference(mws, min_col=2, min_row=7, max_row=18))
    ch.x_axis.number_format = "mmm"
    ch.y_axis.numFmt = '"$"#,##0'
    ch.series[0].graphicalProperties.solidFill = ACCENT
    ch.series[1].graphicalProperties.solidFill = "E0A458"
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.gapWidth = 60
    ws.add_chart(ch, "B13")
    return ws


def start_sheet(wb):
    ws = wb.create_sheet("Start Here", 0)
    K.start_here(ws, "Freelancer Income & Tax Tracker",
                 "Log what comes in and what goes out, see income minus deductible expenses each month, and know how much to set aside for tax. "
                 "Works in Microsoft Excel and Google Sheets.",
                 ["Open Setup. Enter the tax year, your business name, the % of income minus deductible expenses you want to set aside for tax, and an optional monthly income goal.",
                  "Still on Setup: list your clients and check the expense categories. Change the names or deductible % to fit your work.",
                  "Log every payment you receive on Income, and every business cost on Expenses. Dropdowns fill from Setup.",
                  "Use Invoices to track what you have billed. Overdue invoices turn red.",
                  "Check Monthly & Tax for income minus deductible expenses and set-aside per month and per quarter. A loss month shows a negative set-aside (money you no longer need to add). Record tax payments you make there.",
                  "Dashboard and Breakdown update on their own."],
                 ["Google Sheets: upload the file to Google Drive, then open it with Google Sheets (or File > Import).",
                  "Dollar signs are just formatting. To change currency, select the cells and use Format > Number.",
                  "Room for 500 income rows, 500 expense rows and 200 invoices. Start a fresh copy each tax year.", "An expense with no category, or a category not listed on Setup, counts as 0% deductible until you pick one.",
                  "This tracker helps you save for tax. It does not calculate the tax you owe and is not tax advice. Talk to a tax professional about your situation.",
                  "Please do not type over the white formula cells. If something breaks, re-download the original from Etsy (Purchases > Download files)."])


def build(path, sample=False):
    wb = Workbook()
    wb.remove(wb.active)
    setup_sheet(wb, sample)
    income_sheet(wb, sample)
    expense_sheet(wb, sample)
    invoice_sheet(wb, sample)
    monthly_sheet(wb)
    breakdown_sheet(wb)
    dashboard_sheet(wb)
    start_sheet(wb)
    order = ["Start Here", "Dashboard", "Setup", "Income", "Expenses", "Invoices", "Monthly & Tax", "Breakdown"]
    wb._sheets = [wb[n] for n in order]
    if sample:
        for _n, _a in {'Start Here': 'A1:D24', 'Dashboard': 'A1:M30', 'Setup': 'A1:G37', 'Income': 'A1:H32', 'Expenses': 'A1:I32', 'Invoices': 'A1:J16', 'Monthly & Tax': 'A1:K37', 'Breakdown': 'A1:J27'}.items():
            wb[_n].print_area = _a
            wb[_n].page_setup.fitToHeight = 1
            wb[_n].page_setup.fitToWidth = 1
    wb.active = 0
    wb.save(path)
    print(path)


if __name__ == "__main__":
    here = os.path.dirname(os.path.abspath(__file__))
    out = os.path.join(here, "..", "01-freelancer-income-tax-tracker")
    os.makedirs(out, exist_ok=True)
    build(os.path.join(out, "Freelancer-Income-Tax-Tracker.xlsx"))
    samp = sys.argv[1] if len(sys.argv) > 1 else os.path.join(here, "samples")
    os.makedirs(samp, exist_ok=True)
    build(os.path.join(samp, "Freelancer-Income-Tax-Tracker_SAMPLE.xlsx"), sample=True)
