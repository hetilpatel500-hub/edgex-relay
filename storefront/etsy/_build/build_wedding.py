"""Wedding Budget Planner: budget split, costs & payments, guest list/RSVP, checklist."""
import os
import sys
import random
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule, DataBarRule
from xlsxkit import Kit, INK, MUTED

ACCENT = "9A5B6F"
K = Kit(ACCENT)
MONEY = '"$"#,##0.00'
MONEY0 = '"$"#,##0'
DATEF = 'mmm d, yyyy'
CAT0, CAT1 = 8, 27          # 20 category rows on Budget
X0, X1 = 7, 156             # 150 cost rows
G0, G1 = 7, 306             # 300 guest rows
T0 = 7

CATS = [("Venue & rentals", .28), ("Catering & bar", .22), ("Photography & video", .12), ("Attire & beauty", .08),
        ("Flowers & decor", .08), ("Music & entertainment", .06), ("Planner / coordinator", .03), ("Rings", .03),
        ("Cake & desserts", .02), ("Stationery & postage", .02), ("Transportation", .02), ("Favors & gifts", .02),
        ("Officiant & license", .01), ("Buffer for surprises", .01)]

CHECK = [
    ("12+ months", ["Set your total budget and decide who is paying for what", "Draft a rough guest list and head count",
                    "Pick a season or date range", "Book the venue", "Decide if you want a planner or coordinator"]),
    ("9-12 months", ["Book photographer and videographer", "Book caterer (if not with venue)", "Book officiant",
                     "Start looking at attire", "Set up a wedding website (optional)", "Reserve hotel blocks for guests"]),
    ("6-9 months", ["Book music / DJ / band", "Book florist", "Send save-the-dates", "Order attire",
                    "Plan the honeymoon", "Register for gifts"]),
    ("4-6 months", ["Book hair and makeup", "Order invitations", "Book transportation", "Order cake / desserts",
                    "Buy rings", "Plan rehearsal dinner"]),
    ("2-3 months", ["Mail invitations", "First attire fitting", "Finalize menu and bar", "Write or plan vows",
                    "Buy gifts for the wedding party", "Check marriage license requirements"]),
    ("1 month", ["Chase missing RSVPs", "Give final numbers to caterer and venue", "Make the seating plan",
                 "Get the marriage license", "Final attire fitting", "Confirm timings with every vendor"]),
    ("Week of", ["Pack an emergency kit", "Put final payments and tips in labelled envelopes", "Give vendors a day-of contact",
                 "Rehearsal", "Rest and eat well"]),
    ("After", ["Return rentals", "Send thank-you notes", "Change name on documents (if you choose)", "Leave vendor reviews",
               "Preserve the dress / bouquet (optional)"]),
]


def budget_sheet(wb, sample):
    ws = wb.create_sheet("Budget")
    K.sheet_setup(ws, [2, 28, 12, 15, 15, 15, 15, 15, 3], tab=ACCENT)
    K.title(ws, "Budget plan", "Set your total, split it by category, and watch real costs against the plan.", span=7)
    K.label(ws["B5"], "Total budget")
    ws["C5"] = 30000 if sample else 0
    K.body(ws["C5"], MONEY0, inp=True)
    ws.merge_cells("C5:D5")
    K.label(ws["E5"], "Split used")
    ws["F5"] = "=SUM(C{a}:C{b})".format(a=CAT0, b=CAT1)
    K.body(ws["F5"], "0%", bold=True, band=True, align="center")
    ws["G5"] = '=IF(ROUND(F5,4)=1,"Adds up to 100%",IF(F5<1,"Under 100%: "&TEXT(1-F5,"0%")&" unassigned","Over 100% by "&TEXT(F5-1,"0%")))'
    ws["G5"].font = Font(name="Arial", size=9, bold=True, color=ACCENT)
    K.header_row(ws, 7, 2, ["Category", "% of budget", "Planned", "Committed cost", "Paid so far", "Still to pay", "Left in plan"])
    for i in range(CAT1 - CAT0 + 1):
        r = CAT0 + i
        name, pct = CATS[i] if i < len(CATS) else (None, None)
        ws.cell(row=r, column=2, value=name)
        ws.cell(row=r, column=3, value=pct)
        ws.cell(row=r, column=4, value='=IF(B{r}="","",$C$5*C{r})'.format(r=r))
        crit = "Costs!$B${a}:$B${b},$B{r}".format(a=X0, b=X1, r=r)
        ws.cell(row=r, column=5, value='=IF(B{r}="","",SUMIFS(Costs!$F${a}:$F${b},{c}))'.format(r=r, a=X0, b=X1, c=crit))
        ws.cell(row=r, column=6, value='=IF(B{r}="","",SUMIFS(Costs!$G${a}:$G${b},{c}))'.format(r=r, a=X0, b=X1, c=crit))
        ws.cell(row=r, column=7, value='=IF(B{r}="","",E{r}-F{r})'.format(r=r))
        ws.cell(row=r, column=8, value='=IF(B{r}="","",D{r}-E{r})'.format(r=r))
        band = bool(i % 2)
        K.body(ws.cell(row=r, column=2), inp=True)
        K.body(ws.cell(row=r, column=3), "0%", inp=True, align="center")
        for c in range(4, 9):
            K.body(ws.cell(row=r, column=c), MONEY, band=band)
    t = CAT1 + 1
    K.label(ws.cell(row=t, column=2), "Totals", color=INK)
    for c, col in ((4, "D"), (5, "E"), (6, "F"), (7, "G"), (8, "H")):
        ws.cell(row=t, column=c, value="=SUM({c}{a}:{c}{b})".format(c=col, a=CAT0, b=CAT1))
        K.body(ws.cell(row=t, column=c), MONEY, bold=True, band=True)
    ws.conditional_formatting.add("H{a}:H{b}".format(a=CAT0, b=t), FormulaRule(formula=["AND(ISNUMBER(H{a}),H{a}<0)".format(a=CAT0)],
                                  font=Font(color="A33A2B", bold=True), fill=K.fill("FBE3DC")))
    K.note(ws, "B%d" % (t + 2), "The starting split is only a suggestion to get you going. Change the % to match what matters to you. Red = over plan.", italic=True)
    ws.freeze_panes = "A8"
    return ws


def costs_sheet(wb, sample):
    ws = wb.create_sheet("Costs")
    K.sheet_setup(ws, [2, 24, 26, 22, 18, 14, 14, 14, 15, 18, 26], tab=K.light)
    K.title(ws, "Costs & payments", "One row per vendor or purchase. Deposits and balances in one place.", span=10)
    K.header_row(ws, 6, 2, ["Category", "Item", "Vendor", "Contact", "Committed cost", "Paid so far", "Still to pay", "Next payment due", "Status", "Notes"])
    K.list_validation(ws, "B{a}:B{b}".format(a=X0, b=X1), "=Budget!$B${a}:$B${b}".format(a=CAT0, b=CAT1), "Categories come from the Budget tab.")
    for i in range(X1 - X0 + 1):
        r = X0 + i
        for c, fmt in ((2, None), (3, None), (4, None), (5, None), (6, MONEY), (7, MONEY), (9, DATEF), (11, None)):
            K.body(ws.cell(row=r, column=c), fmt, inp=True)
        ws.cell(row=r, column=8, value='=IF(F{r}="","",F{r}-N(G{r}))'.format(r=r))
        ws.cell(row=r, column=10, value=('=IF(F{r}="","",IF(H{r}<=0,"Paid in full",IF(I{r}="","Balance due",'
                                         'IF(I{r}<TODAY(),"Overdue",IF(I{r}-TODAY()<=30,"Due within 30 days","Scheduled")))))').format(r=r))
        K.body(ws.cell(row=r, column=8), MONEY)
        K.body(ws.cell(row=r, column=10), None, align="center")
    rng = "J{a}:J{b}".format(a=X0, b=X1)
    ws.conditional_formatting.add(rng, FormulaRule(formula=['$J{a}="Overdue"'.format(a=X0)], fill=K.fill("FBE3DC"), font=Font(color="A33A2B", bold=True)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=['$J{a}="Due within 30 days"'.format(a=X0)], fill=K.fill("FFF1D6"), font=Font(color="8A5A00", bold=True)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=['$J{a}="Paid in full"'.format(a=X0)], font=Font(color="2F7D6D", bold=True)))
    if sample:
        today = dt.date(2026, 9, 25)
        rows = [("Venue & rentals", "Venue hire", "Willow Barn", 7800, 3900, 60), ("Catering & bar", "Dinner for 110", "Fork & Field", 6600, 1500, 90),
                ("Photography & video", "Photographer, 8 hrs", "Lumen Photo", 3400, 1000, 20), ("Attire & beauty", "Dress & alterations", "Bridal studio", 2100, 2100, None),
                ("Flowers & decor", "Flowers & arch", "Petal House", 2300, 500, -5), ("Music & entertainment", "DJ", "Night Owl DJs", 1600, 400, 45),
                ("Cake & desserts", "3-tier cake", "Sugar Loaf", 650, 150, 100), ("Stationery & postage", "Invitations", "Print shop", 480, 480, None),
                ("Rings", "Two bands", "Jeweler", 1100, 1100, None), ("Transportation", "Guest shuttle", "City Shuttle", 700, 0, 70)]
        for i, (cat, item, v, cost, paid, due) in enumerate(rows):
            r = X0 + i
            ws.cell(row=r, column=2, value=cat)
            ws.cell(row=r, column=3, value=item)
            ws.cell(row=r, column=4, value=v)
            ws.cell(row=r, column=6, value=cost)
            ws.cell(row=r, column=7, value=paid)
            if due is not None:
                ws.cell(row=r, column=9, value=today + dt.timedelta(days=due))
    ws.freeze_panes = "A7"
    return ws


def guests_sheet(wb, sample):
    ws = wb.create_sheet("Guests")
    K.sheet_setup(ws, [2, 26, 12, 9, 26, 11, 11, 12, 11, 16, 20, 9, 18, 11], tab=K.light)
    K.title(ws, "Guest list & RSVPs", "One row per invitation. Put the number of people it covers in Party size.", span=13)
    K.header_row(ws, 6, 2, ["Name(s)", "Side", "Party size", "Email / phone", "Save-the-date", "Invite sent", "RSVP",
                             "Attending", "Meal choice", "Dietary needs", "Table", "Gift", "Thank-you sent"])
    K.list_validation(ws, "C{a}:C{b}".format(a=G0, b=G1), '"Partner 1,Partner 2,Both"')
    K.list_validation(ws, "F{a}:G{b}".format(a=G0, b=G1), '"Yes,No"')
    K.list_validation(ws, "H{a}:H{b}".format(a=G0, b=G1), '"Yes,No,Waiting"')
    K.list_validation(ws, "J{a}:J{b}".format(a=G0, b=G1), '"Option 1,Option 2,Option 3,Kids meal"')
    K.list_validation(ws, "N{a}:N{b}".format(a=G0, b=G1), '"Yes,No"')
    for i in range(G1 - G0 + 1):
        r = G0 + i
        for c in (2, 3, 4, 5, 6, 7, 8, 10, 11, 12, 13, 14):
            K.body(ws.cell(row=r, column=c), "0" if c in (4, 12) else None, inp=True,
                   align="center" if c in (3, 4, 6, 7, 8, 12, 14) else None)
        ws.cell(row=r, column=9, value='=IF(H{r}="Yes",N(D{r}),IF(B{r}="","",0))'.format(r=r))
        K.body(ws.cell(row=r, column=9), "0", align="center")
    rng = "H{a}:H{b}".format(a=G0, b=G1)
    ws.conditional_formatting.add(rng, FormulaRule(formula=['$H{a}="Yes"'.format(a=G0)], font=Font(color="2F7D6D", bold=True)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=['$H{a}="No"'.format(a=G0)], font=Font(color="A33A2B", bold=True)))
    if sample:
        rnd = random.Random(5)
        first = ["Alex", "Sam", "Jordan", "Priya", "Chen", "Maria", "Liam", "Aisha", "Noah", "Emma", "Omar", "Grace", "Leo", "Ivy", "Mateo"]
        last = ["Rivera", "Kim", "Patel", "Nguyen", "Brooks", "Garcia", "Hughes", "Okafor", "Silva", "Novak"]
        for i in range(48):
            r = G0 + i
            n = "%s %s" % (rnd.choice(first), rnd.choice(last))
            size = rnd.choice([1, 2, 2, 2, 3, 4])
            if size == 2:
                n += " & guest"
            ws.cell(row=r, column=2, value=n)
            ws.cell(row=r, column=3, value=rnd.choice(["Partner 1", "Partner 2", "Both"]))
            ws.cell(row=r, column=4, value=size)
            ws.cell(row=r, column=6, value="Yes")
            ws.cell(row=r, column=7, value="Yes")
            rs = rnd.choice(["Yes", "Yes", "Yes", "No", "Waiting"])
            ws.cell(row=r, column=8, value=rs)
            if rs == "Yes":
                ws.cell(row=r, column=10, value=rnd.choice(["Option 1", "Option 2", "Option 3"]))
                ws.cell(row=r, column=12, value=rnd.randint(1, 12))
    ws.freeze_panes = "C7"
    return ws


def checklist_sheet(wb):
    ws = wb.create_sheet("Checklist")
    K.sheet_setup(ws, [2, 16, 58, 10, 16, 28], tab=K.light, landscape=False)
    K.title(ws, "Planning checklist", "A timeline to work from. Mark tasks Done, add your own in the blank rows.", span=5)
    K.label(ws["B5"], "Progress")
    K.header_row(ws, 6, 2, ["When", "Task", "Done?", "Deadline", "Notes"])
    r = T0
    for when, tasks in CHECK:
        for t in tasks + [""]:
            ws.cell(row=r, column=2, value=when)
            ws.cell(row=r, column=3, value=t or None)
            K.body(ws.cell(row=r, column=2), None, bold=True, color=ACCENT, band=True)
            for c in (3, 4, 5, 6):
                K.body(ws.cell(row=r, column=c), DATEF if c == 5 else None, inp=True, align="center" if c == 4 else None)
            r += 1
    last = r - 1
    K.list_validation(ws, "D{a}:D{b}".format(a=T0, b=last), '"Done,Skip"')
    ws["C5"] = '=IFERROR(COUNTIF(D{a}:D{b},"Done")/(COUNTA(C{a}:C{b})-COUNTIF(D{a}:D{b},"Skip")),0)'.format(a=T0, b=last)
    K.body(ws["C5"], "0%", bold=True, band=True)
    ws.conditional_formatting.add("C5", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=ACCENT))
    ws.conditional_formatting.add("C{a}:D{b}".format(a=T0, b=last), FormulaRule(formula=['$D{a}="Done"'.format(a=T0)], font=Font(color="9AA3AA", strike=True)))
    ws.freeze_panes = "A7"
    return ws, last


def dashboard_sheet(wb, sample, last_task):
    ws = wb.create_sheet("Dashboard", 0)
    K.sheet_setup(ws, [2] + [13] * 12, tab=ACCENT)
    K.title(ws, "Our Wedding Dashboard", None, span=12)
    K.label(ws["B5"], "Wedding date")
    ws["D5"] = dt.date(2027, 6, 12) if sample else None
    K.body(ws["D5"], DATEF, inp=True)
    ws.merge_cells("D5:E5")
    K.label(ws["G5"], "Names (optional)")
    ws["I5"] = "Jamie & Taylor" if sample else None
    K.body(ws["I5"], inp=True)
    ws.merge_cells("I5:K5")
    K.kpi(ws, 7, 2, "Days to go", '=IF(D5="","Add date",MAX(0,D5-TODAY()))', "0", 2)
    K.kpi(ws, 7, 4, "Total budget", "=Budget!C5", MONEY0, 3)
    K.kpi(ws, 7, 7, "Committed so far", "=Budget!E%d" % (CAT1 + 1), MONEY0, 3)
    K.kpi(ws, 7, 10, "Left to commit", "=Budget!C5-Budget!E%d" % (CAT1 + 1), MONEY0, 3)
    K.kpi(ws, 10, 2, "Paid so far", "=Budget!F%d" % (CAT1 + 1), MONEY0, 2)
    K.kpi(ws, 10, 4, "Still to pay", "=Budget!G%d" % (CAT1 + 1), MONEY0, 3)
    K.kpi(ws, 10, 7, "Payments overdue", '=COUNTIF(Costs!J{a}:J{b},"Overdue")'.format(a=X0, b=X1), "0", 3)
    K.kpi(ws, 10, 10, "Cost per guest", '=IF(N(E15)=0,"-",Budget!E{t}/E15)'.format(t=CAT1 + 1), MONEY0, 3)
    K.section(ws["B13"], "Guests (people, not invitations)")
    g = dict(a=G0, b=G1)
    K.kpi(ws, 14, 2, "Invited", "=SUM(Guests!D{a}:D{b})".format(**g), "0", 3)
    K.kpi(ws, 14, 5, "Attending", "=SUM(Guests!I{a}:I{b})".format(**g), "0", 3)
    K.kpi(ws, 14, 8, "Declined", '=SUMIFS(Guests!D{a}:D{b},Guests!H{a}:H{b},"No")'.format(**g), "0", 3)
    K.kpi(ws, 14, 11, "Still waiting", "=B15-E15-H15", "0", 2)
    K.section(ws["B17"], "Checklist")
    ws["B18"] = "Tasks done"
    ws["B18"].font = Font(name="Arial", size=8, bold=True, color=MUTED)
    ws["D18"] = "=Checklist!C5"
    ws["D18"].number_format = "0%"
    ws["D18"].font = Font(name="Georgia", size=16, bold=True, color=INK)
    ch = BarChart()
    ch.type = "bar"
    ch.title = "Planned vs committed by category"
    ch.height, ch.width = 10, 24
    b = wb["Budget"]
    ch.add_data(Reference(b, min_col=4, min_row=7, max_row=7 + len(CATS)), titles_from_data=True)
    ch.add_data(Reference(b, min_col=5, min_row=7, max_row=7 + len(CATS)), titles_from_data=True)
    ch.set_categories(Reference(b, min_col=2, min_row=8, max_row=7 + len(CATS)))
    ch.series[0].graphicalProperties.solidFill = "D9B8C2"
    ch.series[1].graphicalProperties.solidFill = ACCENT
    ch.x_axis.scaling.orientation = "maxMin"
    ch.y_axis.numFmt = '"$"#,##0'
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.gapWidth = 40
    ws.add_chart(ch, "B20")
    return ws


def start_sheet(wb):
    ws = wb.create_sheet("Start Here", 0)
    K.start_here(ws, "Wedding Budget Planner",
                 "Plan your budget, track every vendor payment, manage RSVPs and tick off your to-dos, all in one file. "
                 "Works in Microsoft Excel and Google Sheets.",
                 ["On Dashboard, type your wedding date (and names if you like).",
                  "On Budget, type your total budget. Adjust the % split by category until it says 'Adds up to 100%'. Rename or add categories in the blank rows.",
                  "On Costs, add each vendor or purchase as you book it: committed cost, what you have paid, and when the next payment is due. Status updates by itself.",
                  "On Guests, add each invitation and its party size. Update RSVPs as they come in. Attending counts fill in automatically.",
                  "Work through Checklist and mark tasks Done. Progress shows on the Dashboard.",
                  "Dashboard shows days to go, money committed and paid, overdue payments, cost per guest and RSVPs."],
                 ["Google Sheets: upload the file to Google Drive, then open it with Google Sheets.",
                  "Dollar signs are only formatting. To change currency, select the cells and use Format > Number.",
                  "Room for 20 budget categories, 150 costs and 300 invitations.",
                  "The starting % split is a suggestion, not a rule. Every wedding is different.",
                  "Please do not type over the white formula cells. If something breaks, re-download the original from Etsy (Purchases > Download files)."])


def build(path, sample=False):
    wb = Workbook()
    wb.remove(wb.active)
    budget_sheet(wb, sample)
    costs_sheet(wb, sample)
    guests_sheet(wb, sample)
    _, last = checklist_sheet(wb)
    dashboard_sheet(wb, sample, last)
    start_sheet(wb)
    order = ["Start Here", "Dashboard", "Budget", "Costs", "Guests", "Checklist"]
    wb._sheets = [wb[n] for n in order]
    wb.active = 0
    if sample:
        ws = wb["Checklist"]
        for r in range(T0, T0 + 14):
            if ws.cell(row=r, column=3).value:
                ws.cell(row=r, column=4, value="Done")
    if sample:
        for _n, _a in {'Start Here': 'A1:D22', 'Dashboard': 'A1:M42', 'Budget': 'A1:H31', 'Costs': 'A1:K20', 'Guests': 'A1:N32', 'Checklist': 'A1:F40'}.items():
            wb[_n].print_area = _a
            wb[_n].page_setup.fitToHeight = 1
            wb[_n].page_setup.fitToWidth = 1
    wb.save(path)
    print(path)


if __name__ == "__main__":
    here = os.path.dirname(os.path.abspath(__file__))
    out = os.path.join(here, "..", "02-wedding-budget-planner")
    os.makedirs(out, exist_ok=True)
    build(os.path.join(out, "Wedding-Budget-Planner.xlsx"))
    samp = sys.argv[1] if len(sys.argv) > 1 else os.path.join(here, "samples")
    os.makedirs(samp, exist_ok=True)
    build(os.path.join(samp, "Wedding-Budget-Planner_SAMPLE.xlsx"), sample=True)
