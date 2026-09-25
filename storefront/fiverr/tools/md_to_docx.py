"""Turn a simple Markdown deliverable into a clean Word (.docx) file.

Used for Fiverr Gigs 1 and 2 (the gigs promise "a Word file").
Handles: # / ## / ### headings, paragraphs, - bullets, 1. numbered items,
**bold**, *italic*, > note blocks, --- rules, and | tables |.

Usage: python3 md_to_docx.py in.md out.docx
"""
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

NAVY = RGBColor(0x14, 0x21, 0x3D)


def add_runs(p, text):
    for part in re.split(r"(\*\*.+?\*\*|(?<![\w*])\*[^*\s][^*]*?\*(?!\w))", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            p.add_run(part[1:-1]).italic = True
        else:
            p.add_run(re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)", r"\1 (\2)", part).replace("`", ""))


def shade(p, hex_fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    pPr.append(shd)


def convert(src, dst):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)
    for lvl in (1, 2, 3):
        hs = doc.styles[f"Heading {lvl}"]
        hs.font.color.rgb, hs.font.name = NAVY, "Calibri"
    L = open(src, encoding="utf-8").read().split("\n")
    i = 0
    while i < len(L):
        s = L[i].strip()
        if s.startswith(">"):
            buf = []
            while i < len(L) and L[i].strip().startswith(">"):
                buf.append(L[i].strip()[1:].strip())
                i += 1
            p = doc.add_paragraph()
            add_runs(p, " ".join(buf))
            shade(p, "FDEDE8")
            continue
        if s.startswith("|"):
            rows = []
            while i < len(L) and L[i].strip().startswith("|"):
                r = [c.strip() for c in L[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"-+", c) for c in r if c):
                    rows.append(r)
                i += 1
            t = doc.add_table(rows=len(rows), cols=len(rows[0]))
            t.style, t.alignment = "Light Grid Accent 1", WD_TABLE_ALIGNMENT.CENTER
            for a, r in enumerate(rows):
                for b, c in enumerate(r):
                    cell = t.cell(a, b)
                    cell.text = ""
                    add_runs(cell.paragraphs[0], c)
            continue
        m = re.match(r"^(#{1,3}) (.*)", s)
        if m:
            doc.add_heading(m.group(2).replace("**", ""), level=len(m.group(1)))
        elif s == "---":
            doc.add_paragraph("")
        elif re.match(r"^[-*] ", s) or re.match(r"^\d+\. ", s):
            buf = [s]
            i += 1
            while i < len(L) and L[i].startswith("  ") and L[i].strip():
                buf.append(L[i].strip())
                i += 1
            txt = " ".join(buf)
            numbered = bool(re.match(r"^\d+\. ", txt))
            p = doc.add_paragraph(style="List Number" if numbered else "List Bullet")
            add_runs(p, txt.split(". ", 1)[1] if numbered else txt[2:])
            continue
        elif s:
            buf = [s]
            i += 1
            while i < len(L) and L[i].strip() and not re.match(r"^(#|[-*] |\d+\. |\||>|---)", L[i].strip()):
                if L[i].strip().startswith("**"):
                    break
                buf.append(L[i].strip())
                i += 1
            add_runs(doc.add_paragraph(), " ".join(buf))
            continue
        i += 1
    doc.save(dst)
    print("wrote", dst)


if __name__ == "__main__":
    convert(sys.argv[1], sys.argv[2])
