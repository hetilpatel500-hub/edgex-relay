"""Build a clean, editable 16:9 pitch deck (.pptx) from a JSON spec.

Used for Fiverr Gig 3. No watermark, no paid tools: python-pptx only.
Charts are native PowerPoint charts, so buyers can edit the numbers.

Usage:
    python3 build_deck.py spec.json out.pptx
    soffice --headless --convert-to pdf out.pptx   # PDF copy for delivery

Spec format (see samples/gig3-hearthloop-spec.json for a full example):
{
  "colors": {"primary": "#14213D", "accent": "#F2785C", "bg": "#FAF7F2", "text": "#1F2430"},
  "font": "Calibri",
  "footer": "Hearthloop | Seed 2026",
  "slides": [
    {"type": "title", "title": "...", "subtitle": "...", "note": "..."},
    {"type": "bullets", "title": "...", "bullets": ["..."], "notes": "..."},
    {"type": "two_col", "title": "...", "left": {"heading": "...", "bullets": []}, "right": {...}},
    {"type": "steps", "title": "...", "steps": [{"heading": "...", "text": "..."}]},
    {"type": "market", "title": "...", "rings": [{"label": "TAM", "value": "...", "text": "..."}], "source": "..."},
    {"type": "chart", "title": "...", "categories": [], "series": {"name": [values]}, "caption": "..."},
    {"type": "table", "title": "...", "columns": [], "rows": [[]], "source": "..."},
    {"type": "team", "title": "...", "people": [{"name": "...", "role": "...", "text": "..."}]},
    {"type": "ask", "title": "...", "amount": "...", "uses": [{"label": "...", "pct": 40}], "contact": "..."}
  ]
}
Every slide type accepts "notes" (speaker notes).
"""
import json
import sys

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

W, H = Inches(13.333), Inches(7.5)
M = Inches(0.7)


def rgb(h):
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def mix(h1, h2, t):
    a, b = rgb(h1), rgb(h2)
    return RGBColor(*[int(a[i] + (b[i] - a[i]) * t) for i in range(3)])


class Deck:
    def __init__(self, spec):
        c = spec.get("colors", {})
        self.P = c.get("primary", "#14213D")
        self.A = c.get("accent", "#F2785C")
        self.BG = c.get("bg", "#FAF7F2")
        self.T = c.get("text", "#1F2430")
        self.font = spec.get("font", "Calibri")
        self.footer = spec.get("footer", "")
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.blank = self.prs.slide_layouts[6]
        self.n = 0

    # ---------- primitives ----------
    def rect(self, s, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE, line=None):
        r = s.shapes.add_shape(shape, x, y, w, h)
        r.fill.solid()
        r.fill.fore_color.rgb = fill if isinstance(fill, RGBColor) else rgb(fill)
        if line:
            r.line.color.rgb = rgb(line)
            r.line.width = Pt(1)
        else:
            r.line.fill.background()
        r.shadow.inherit = False
        return r

    def text(self, s, x, y, w, h, txt, size=18, bold=False, color=None,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False):
        tb = s.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.05)
        lines = txt if isinstance(txt, list) else [txt]
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run()
            r.text = line
            f = r.font
            f.name, f.size, f.bold, f.italic = self.font, Pt(size), bold, italic
            f.color.rgb = rgb(color or self.T)
        return tb

    def bullets(self, s, x, y, w, h, items, size=20, color=None, gap=10):
        tb = s.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(gap)
            dot = p.add_run()
            dot.text = "■  "
            dot.font.size, dot.font.color.rgb, dot.font.name = Pt(size * 0.6), rgb(self.A), self.font
            r = p.add_run()
            r.text = item
            r.font.size, r.font.name = Pt(size), self.font
            r.font.color.rgb = rgb(color or self.T)
        return tb

    def base(self, title, dark=False, notes=None):
        s = self.prs.slides.add_slide(self.blank)
        self.n += 1
        self.rect(s, 0, 0, W, H, self.P if dark else self.BG)
        if not dark:
            self.rect(s, M, Inches(0.55), Inches(0.9), Inches(0.09), self.A)
            self.text(s, M, Inches(0.7), W - 2 * M, Inches(1.0), title, size=34, bold=True, color=self.P)
            foot = f"{self.footer}    {self.n}" if self.footer else str(self.n)
            self.text(s, M, H - Inches(0.55), W - 2 * M, Inches(0.35), foot, size=11,
                      color="#8A8F9C", align=PP_ALIGN.RIGHT)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    # ---------- slide types ----------
    def title(self, d):
        s = self.base("", dark=True, notes=d.get("notes"))
        self.rect(s, M, Inches(2.2), Inches(1.1), Inches(0.12), self.A)
        self.text(s, M, Inches(2.5), Inches(11), Inches(1.4), d["title"], size=60, bold=True, color=self.BG)
        self.text(s, M, Inches(3.9), Inches(10.5), Inches(1.2), d.get("subtitle", ""), size=26,
                  color="#C9CFDB")
        if d.get("note"):
            self.text(s, M, H - Inches(1.3), Inches(11.5), Inches(0.8), d["note"], size=14, color="#9AA3B5",
                      italic=True)

    def bullets_slide(self, d):
        s = self.base(d["title"], notes=d.get("notes"))
        has_side = bool(d.get("callout"))
        w = Inches(7.6) if has_side else W - 2 * M
        self.bullets(s, M, Inches(2.0), w, Inches(4.6), d["bullets"], size=d.get("size", 22))
        if has_side:
            x = M + Inches(8.1)
            self.rect(s, x, Inches(2.0), Inches(3.8), Inches(3.9), self.P, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            self.text(s, x + Inches(0.35), Inches(2.3), Inches(3.1), Inches(3.3), d["callout"], size=22,
                      bold=True, color=self.BG, anchor=MSO_ANCHOR.MIDDLE)

    def two_col(self, d):
        s = self.base(d["title"], notes=d.get("notes"))
        cw = (W - 2 * M - Inches(0.5)) / 2
        for i, col in enumerate([d["left"], d["right"]]):
            x = M + i * (cw + Inches(0.5))
            fill = self.P if i == 1 else mix(self.BG, self.P, 0.08)
            fg = self.BG if i == 1 else self.T
            self.rect(s, x, Inches(2.0), cw, Inches(4.3), fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            self.text(s, x + Inches(0.45), Inches(2.35), cw - Inches(0.9), Inches(0.7), col["heading"],
                      size=28, bold=True, color=self.A if i == 1 else self.P)
            self.bullets(s, x + Inches(0.45), Inches(3.2), cw - Inches(0.9), Inches(2.9), col["bullets"],
                         size=21, color=fg if i == 1 else None, gap=12)

    def steps(self, d):
        s = self.base(d["title"], notes=d.get("notes"))
        k = len(d["steps"])
        gap = Inches(0.35)
        cw = (W - 2 * M - gap * (k - 1)) / k
        for i, st in enumerate(d["steps"]):
            x = M + i * (cw + gap)
            self.rect(s, x, Inches(2.1), cw, Inches(3.6), "#FFFFFF", shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                      line="#E3DED5")
            c = self.rect(s, x + Inches(0.35), Inches(2.4), Inches(0.75), Inches(0.75), self.A, shape=MSO_SHAPE.OVAL)
            c.text_frame.text = str(i + 1)
            p = c.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            p.runs[0].font.size, p.runs[0].font.bold, p.runs[0].font.name = Pt(24), True, self.font
            p.runs[0].font.color.rgb = rgb("#FFFFFF")
            self.text(s, x + Inches(0.35), Inches(3.35), cw - Inches(0.7), Inches(0.7), st["heading"], size=24,
                      bold=True, color=self.P)
            self.text(s, x + Inches(0.35), Inches(4.05), cw - Inches(0.7), Inches(1.5), st["text"], size=19)

    def market(self, d):
        s = self.base(d["title"], notes=d.get("notes"))
        rings = d["rings"]
        cx, cy, R = M + Inches(2.9), Inches(3.95), Inches(2.05)
        for i, rg in enumerate(rings):
            r = int(R * (1 - i * 0.3))
            fill = mix(self.P, self.BG, 0.75 - i * 0.3) if i < len(rings) - 1 else rgb(self.A)
            self.rect(s, Emu(int(cx - r)), Emu(int(cy + R - 2 * r)), Emu(2 * r), Emu(2 * r), fill, shape=MSO_SHAPE.OVAL)
            self.text(s, Emu(int(cx - r)), Emu(int(cy + R - 2 * r + Inches(0.12))), Emu(2 * r), Inches(0.5),
                      rg["label"], size=16, bold=True, color="#FFFFFF" if i else self.P, align=PP_ALIGN.CENTER)
        x = M + Inches(6.4)
        y = Inches(2.1)
        for rg in rings:
            self.text(s, x, y, Inches(5.5), Inches(0.5), f'{rg["label"]}  {rg["value"]}', size=24, bold=True,
                      color=self.P)
            self.text(s, x, y + Inches(0.5), Inches(5.5), Inches(0.8), rg["text"], size=16)
            y += Inches(1.45)
        if d.get("source"):
            self.text(s, M, H - Inches(1.0), W - 2 * M, Inches(0.4), d["source"], size=11, color="#6B7180",
                      italic=True)

    def chart(self, d):
        s = self.base(d["title"], notes=d.get("notes"))
        cd = CategoryChartData()
        cd.categories = d["categories"]
        for name, vals in d["series"].items():
            cd.add_series(name, vals)
        gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, M, Inches(1.9), Inches(8.2), Inches(4.8), cd)
        ch = gf.chart
        ch.has_legend = len(d["series"]) > 1
        if ch.has_legend:
            ch.legend.position, ch.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
        ch.font.size, ch.font.name = Pt(13), self.font
        ch.value_axis.has_major_gridlines = True
        ch.value_axis.major_gridlines.format.line.color.rgb = rgb("#E3DED5")
        ch.value_axis.format.line.fill.background()
        cols = [self.P, self.A, "#8A8F9C"]
        for i, ser in enumerate(ch.series):
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = rgb(cols[i % 3])
        ch.plots[0].gap_width = 60
        x = M + Inches(8.7)
        self.rect(s, x, Inches(2.1), Inches(3.2), Inches(4.2), self.P, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        self.text(s, x + Inches(0.3), Inches(2.4), Inches(2.6), Inches(3.6), d.get("caption", ""), size=19,
                  color=self.BG, anchor=MSO_ANCHOR.MIDDLE)

    def table(self, d):
        s = self.base(d["title"], notes=d.get("notes"))
        rows, cols = len(d["rows"]) + 1, len(d["columns"])
        tw = W - 2 * M
        gt = s.shapes.add_table(rows, cols, M, Inches(1.95), tw, Inches(0.6) * rows)
        t = gt.table
        for j, name in enumerate(d["columns"]):
            t.columns[j].width = int(tw / cols)
        for i in range(rows):
            for j in range(cols):
                cell = t.cell(i, j)
                val = d["columns"][j] if i == 0 else d["rows"][i - 1][j]
                cell.text = str(val)
                cell.fill.solid()
                hi = i > 0 and d.get("highlight_row") == i - 1
                cell.fill.fore_color.rgb = (rgb(self.P) if i == 0 else
                                            mix(self.A, "#FFFFFF", 0.8) if hi else
                                            rgb("#FFFFFF") if i % 2 else mix(self.BG, self.P, 0.05))
                cell.margin_left = cell.margin_right = Inches(0.12)
                for p in cell.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size, r.font.name = Pt(15 if i else 16), self.font
                        r.font.bold = i == 0 or j == 0
                        r.font.color.rgb = rgb("#FFFFFF" if i == 0 else self.T)
        if d.get("source"):
            self.text(s, M, H - Inches(1.0), W - 2 * M, Inches(0.4), d["source"], size=11, color="#6B7180",
                      italic=True)

    def team(self, d):
        s = self.base(d["title"], notes=d.get("notes"))
        k = len(d["people"])
        gap = Inches(0.35)
        cw = (W - 2 * M - gap * (k - 1)) / k
        for i, pp in enumerate(d["people"]):
            x = M + i * (cw + gap)
            self.rect(s, x, Inches(2.0), cw, Inches(4.5), "#FFFFFF", shape=MSO_SHAPE.ROUNDED_RECTANGLE, line="#E3DED5")
            self.rect(s, x + (cw - Inches(1.3)) / 2, Inches(2.3), Inches(1.3), Inches(1.3),
                      mix(self.BG, self.P, 0.15), shape=MSO_SHAPE.OVAL)
            self.text(s, x + Inches(0.2), Inches(3.8), cw - Inches(0.4), Inches(0.5), pp["name"], size=20,
                      bold=True, color=self.P, align=PP_ALIGN.CENTER)
            self.text(s, x + Inches(0.2), Inches(4.3), cw - Inches(0.4), Inches(0.4), pp["role"], size=15,
                      color=self.A, bold=True, align=PP_ALIGN.CENTER)
            self.text(s, x + Inches(0.3), Inches(4.8), cw - Inches(0.6), Inches(1.6), pp.get("text", ""), size=14,
                      align=PP_ALIGN.CENTER)

    def ask(self, d):
        s = self.base("", dark=True, notes=d.get("notes"))
        self.text(s, M, Inches(0.8), Inches(6), Inches(0.8), d["title"], size=30, bold=True, color="#C9CFDB")
        self.text(s, M, Inches(1.6), Inches(6), Inches(1.4), d["amount"], size=66, bold=True, color=self.BG)
        self.rect(s, M, Inches(3.15), Inches(1.1), Inches(0.12), self.A)
        if d.get("contact"):
            self.text(s, M, H - Inches(1.5), Inches(6), Inches(1.0), d["contact"], size=16, color="#C9CFDB")
        x, y, bw = M + Inches(6.6), Inches(1.4), Inches(5.3)
        self.text(s, x, y - Inches(0.1), bw, Inches(0.5), "Use of funds", size=20, bold=True, color=self.A)
        y += Inches(0.6)
        for u in d["uses"]:
            self.text(s, x, y, bw, Inches(0.4), f'{u["label"]}  {u["pct"]}%', size=17, color=self.BG)
            self.rect(s, x, y + Inches(0.45), bw, Inches(0.18), "#2A3A5E")
            self.rect(s, x, y + Inches(0.45), int(bw * u["pct"] / 100), Inches(0.18), self.A)
            y += Inches(1.0)

    def build(self, spec, out):
        kinds = {"title": self.title, "bullets": self.bullets_slide, "two_col": self.two_col, "steps": self.steps,
                 "market": self.market, "chart": self.chart, "table": self.table, "team": self.team, "ask": self.ask}
        for d in spec["slides"]:
            kinds[d["type"]](d)
        self.prs.save(out)


if __name__ == "__main__":
    spec = json.load(open(sys.argv[1]))
    Deck(spec).build(spec, sys.argv[2])
    print("wrote", sys.argv[2])
